#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把 D&D 5e 中文人物卡「装备」工作表解析成统一词条库。

输入：card/grids/装备.json   —— xlsx 导出的网格 JSON，键 "行,列"（1 起，列 1 = A）
      card/米瑞尔.xlsx        —— **只读**打开，仅用于恢复网格导出丢失的名称（见下）
输出：dnd-data/equipment.json —— 词条数组，格式见 dnd-data/SCHEMA.md（type = equipment）

本脚本绝不写入 xlsx。

「装备」表是多分块并排表，每个子表有自己的表头 / 列序。脚本按列区间切块、逐块按自己的表头解析：

    A–D   物品 | 特性 | 价格 | 重量          数据 2–95    → 冒险用品 / 弹药 / 法器
    L–Q   工具 | 属性 | 操作 | 制造 | 价格 | 重量  数据 2–42    → 工具
    AQ–AU 护甲名称 | AC | 敏捷加值 | 重量 | 属性    数据 21–53   → 护甲 / 其它(AC来源)
    AI–AP 名称 | 伤害 | 伤害类型 | 词条 | 精通 | 重量 | 价格 | 属性  数据 3–77 → 武器

其余区间是镜像副本、空孪生表、下拉清单、角色自身计算区或查表暂存区，一律不收（原因写在报告里）。

名称丢失
--------
网格导出只含**缓存有值**的单元格。有些名称格是形如 ``=IF(主要!$BG$29="O","巨灵流光","")`` 的公式，
角色没勾选该选项时 Excel 缓存为空串，网格里就整格消失（AQ28 即此情况，与 tools/extract_feats.py 记录的问题同源）。
脚本只读打开 xlsx，把这类公式里的中文字面量取出来当名称（确定性读取，不是猜测）；
打不开 xlsx 或加 --no-xlsx 时该行跳过并在报告中说明。

用法：
    python tools/extract_equipment.py              # 写文件 + 打印解析报告
    python tools/extract_equipment.py --dry-run    # 只打印报告
    python tools/extract_equipment.py --no-xlsx    # 不读 xlsx（丢失公式名）
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
GRID_PATH = ROOT / "card" / "grids" / "装备.json"
XLSX_PATH = ROOT / "card" / "米瑞尔.xlsx"
OUT_PATH = ROOT / "dnd-data" / "equipment.json"
SHEET_NAME = "装备"

# 占位符：不当内容（SCHEMA 硬性约束 4）
PLACEHOLDER_TOKENS = {
    "", "-", "—", "――", "–", "－", "0", "0.0",
    "#VALUE!", "#N/A", "#REF!", "#NAME?", "#DIV/0!", "N/A", "√", "OK",
}


# --------------------------------------------------------------------------- #
# 网格读取
# --------------------------------------------------------------------------- #
def col_index(letters: str) -> int:
    """'A' -> 1, 'AG' -> 33"""
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def col_name(index: int) -> str:
    """1 -> 'A', 33 -> 'AG'"""
    s = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        s = chr(65 + rem) + s
    return s


def is_placeholder(value: str) -> bool:
    return value.strip() in PLACEHOLDER_TOKENS


class Grid:
    def __init__(self, path: Path) -> None:
        raw = json.loads(path.read_text(encoding="utf-8"))
        self.sheet = raw.get("sheet")
        self.rows = int(raw.get("rows", 0))
        self.cols = int(raw.get("cols", 0))
        self.cells: dict[tuple[int, int], str] = {}
        for key, value in raw["cells"].items():
            row, col = key.split(",")
            self.cells[(int(row), int(col))] = value if isinstance(value, str) else str(value)

    def raw(self, row: int, col) -> str:
        """原始单元格值（不 trim、不判占位符）。col 可以是 'A' 或 1。"""
        idx = col_index(col) if isinstance(col, str) else col
        return self.cells.get((row, idx), "")

    def get(self, row: int, col, keep_zero: bool = False) -> str:
        """取内容：占位符（含 '0'）返回空串；keep_zero=True 时保留 '0'。"""
        value = self.raw(row, col).strip()
        if keep_zero:
            return "" if (value in PLACEHOLDER_TOKENS and value != "0") else value
        return "" if is_placeholder(value) else value

    def has(self, row: int, col) -> bool:
        return bool(self.get(row, col))


class NameRecovery:
    """只读打开 xlsx，恢复「公式 + 缓存为空」导致网格缺失的名称格。"""

    def __init__(self, xlsx: Path, sheet: str) -> None:
        self.recovered: dict[tuple[int, str], str] = {}
        self.note = ""
        if not xlsx.exists():
            self.note = f"{xlsx} 不存在，未恢复公式名称"
            return
        try:
            import openpyxl
        except ImportError:
            self.note = "未安装 openpyxl，未恢复公式名称"
            return
        try:
            wb = openpyxl.load_workbook(xlsx, data_only=False, read_only=True)
            cached_wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
            if sheet not in wb.sheetnames:
                self.note = f"xlsx 中没有工作表 {sheet}"
                return
            cached: dict[tuple[int, int], object] = {}
            for r_i, row in enumerate(cached_wb[sheet].iter_rows(), start=1):
                for c_i, cell in enumerate(row, start=1):
                    if cell.value not in (None, ""):
                        cached[(r_i, c_i)] = cell.value
            for r_i, row in enumerate(wb[sheet].iter_rows(), start=1):
                for c_i, cell in enumerate(row, start=1):
                    value = cell.value
                    if not (isinstance(value, str) and value.startswith("=")):
                        continue
                    if (r_i, c_i) in cached:
                        continue                      # 缓存有值 → 网格里本来就有
                    literals = [s for s in re.findall(r'"([^"]*)"', value)
                                if s and s != "O" and re.search(r"[\u4e00-\u9fff]", s)]
                    if not literals:
                        continue
                    self.recovered[(r_i, col_name(c_i))] = max(literals, key=len)
            wb.close()
            cached_wb.close()
            self.note = f"从 xlsx 公式恢复 {len(self.recovered)} 个名称格"
        except Exception as exc:                       # 只读失败不影响主流程
            self.note = f"读取 xlsx 失败（{exc.__class__.__name__}: {exc}），未恢复公式名称"

    def name_for(self, row: int, col: str) -> str:
        return self.recovered.get((row, col.upper()), "")


# --------------------------------------------------------------------------- #
# 词条构造
# --------------------------------------------------------------------------- #
class Builder:
    def __init__(self) -> None:
        self.entries: list[dict] = []
        self.used_ids: dict[str, int] = {}
        self.first_seen: dict[str, int] = {}
        self.collisions: list[tuple[str, int, int]] = []

    def _unique_id(self, name: str, row: int) -> str:
        base = f"equipment:{name}"
        if base not in self.used_ids:
            self.used_ids[base] = 1
            self.first_seen[name] = row
            return base
        self.used_ids[base] += 1
        self.collisions.append((name, self.first_seen[name], row))
        return f"{base}#{self.used_ids[base]}"

    def add(self, name, category, row, *, text="", fields=None, tags=None, source="", en="") -> dict:
        clean_fields = {}
        for k, v in (fields or {}).items():
            v = (v or "").strip()
            # 保留 '0'：调用方已用 Grid.get(keep_zero=False) 把「0 = 无」的字段（价格/重量等）
            # 清成空串，能走到这里的 '0' 是有意义的值（如重甲的敏捷加值 0）。
            if v and (not is_placeholder(v) or v == "0"):
                clean_fields[k] = v
        entry = {
            "id": self._unique_id(name.strip(), row),
            "type": "equipment",
            "name": name.strip(),
            "en": en,
            "category": category,
            "tags": [t for t in (tags or []) if t],
            "summary": "",
            "text": text.strip() if text else "",
            "fields": clean_fields,
            "source": source,
            "cardRef": {"sheet": SHEET_NAME, "row": row},
        }
        self.entries.append(entry)
        return entry


# --------------------------------------------------------------------------- #
# 各分块解析
# --------------------------------------------------------------------------- #
AMMO_NAMES = {"箭矢", "弩矢", "枪械子弹", "投石索子弹", "吹矢"}
FOCUS_NAMES = {
    "水晶", "法球", "权杖", "法杖（也视作长棍）", "魔杖",                 # 奥术法器
    "槲寄生枝条", "木质法杖（也视作长棍）", "紫衫魔杖",                    # 德鲁伊法器
    "护符（佩戴或手持）", "纹章（挂载到织物或盾牌上）", "圣物匣（手持）",  # 圣徽
}
FOCUS_TAGS = {
    "水晶": "奥术法器", "法球": "奥术法器", "权杖": "奥术法器",
    "法杖（也视作长棍）": "奥术法器", "魔杖": "奥术法器",
    "槲寄生枝条": "德鲁伊法器", "木质法杖（也视作长棍）": "德鲁伊法器", "紫衫魔杖": "德鲁伊法器",
    "护符（佩戴或手持）": "圣徽", "纹章（挂载到织物或盾牌上）": "圣徽", "圣物匣（手持）": "圣徽",
}


def parse_main(g: Grid, b: Builder, rep: dict, rec: NameRecovery) -> None:
    """A–D 块：物品 | 特性 | 价格 | 重量（表头第 1 行）。"""
    rows = []
    for r in range(2, g.rows + 1):
        name = g.get(r, "A") or rec.name_for(r, "A")
        price, weight, text = g.get(r, "C"), g.get(r, "D"), g.get(r, "B")
        if not name:
            if price or weight:
                rep["skipped_rows"].append((f"A–D 第 {r} 行", "(空)", "有价格/重量但无名称"))
            continue
        if not (price or weight):
            rep["skipped_rows"].append((f"A–D 第 {r} 行", name, "仅有名称，无价格/重量（表尾残留）"))
            continue
        if name in AMMO_NAMES:
            category, tags = "弹药", []
        elif name in FOCUS_NAMES:
            category, tags = "法器", [FOCUS_TAGS[name]]
        else:
            category, tags = "冒险用品", []
        b.add(name, category, r, text=text, fields={"价格": price, "重量": weight}, tags=tags)
        rows.append(r)
    rep["blocks"].append({
        "name": "主表（冒险用品 / 弹药 / 法器）",
        "cols": "A–D", "header_row": 1, "header": "物品 | 特性 | 价格 | 重量",
        "rows": f"{min(rows)}–{max(rows)}", "count": len(rows),
        "note": "价格、重量列序与 E–H 镜像表相反；弹药 5 条（行 4–8）、法器 11 条（行 10–14 / 40–42 / 49–51）",
    })


def parse_tools(g: Grid, b: Builder, rep: dict, rec: NameRecovery) -> None:
    """L–Q 块：工具 | 属性 | 操作 | 制造 | 价格 | 重量（表头第 1 行）。
    同区间的 K 列是同一份工具清单整体上移一行的辅助列（K[r] == L[r+1]），忽略。"""
    rows = []
    for r in range(2, g.rows + 1):
        name = g.get(r, "L") or rec.name_for(r, "L")
        attr, operate = g.get(r, "M"), g.get(r, "N")
        craft, price, weight = g.get(r, "O"), g.get(r, "P"), g.get(r, "Q")
        if not name:
            if attr or operate or craft or price or weight:
                rep["skipped_rows"].append((f"L–Q 第 {r} 行", "(空)", "有工具数据但无名称"))
            continue
        if not (attr or operate or craft or price or weight):
            continue
        b.add(name, "工具", r, text=operate,
              fields={"属性": attr, "操作": operate, "制造": craft, "价格": price, "重量": weight})
        rows.append(r)
    rep["blocks"].append({
        "name": "工具表", "cols": "L–Q", "header_row": 1,
        "header": "工具 | 属性 | 操作 | 制造 | 价格 | 重量",
        "rows": f"{min(rows)}–{max(rows)}", "count": len(rows),
        "note": "表头行缺 K 列（K 是辅助列）；行 41–42 原表把整段描述写在「操作」列、DC 任务写在「制造」列，按列原样采集",
    })


def parse_armor(g: Grid, b: Builder, rep: dict, rec: NameRecovery) -> None:
    """AQ–AU 块：护甲名称 | AC | 敏捷加值 | 重量 | 属性（表头第 20 行）。
    第 36/41/48 行是 AQ='——' 且 AC=10 的分组分隔行，跳过；
    三个分隔行之后的三组按冒险手册惯例推断为 轻甲 / 中甲 / 重甲（成员与 PHB 完全一致）。"""
    end = max((r for r in range(21, g.rows + 1) if g.get(r, "AR")), default=0)
    group_names = ["轻甲", "中甲", "重甲"]
    group_seq = iter(group_names)
    group = ""
    rows, dividers, healed = [], [], []
    for r in range(21, end + 1):
        raw_name = g.raw(r, "AQ").strip()
        if raw_name == "——":
            dividers.append(r)
            group = next(group_seq, "")
            continue
        name = g.get(r, "AQ")
        note = g.get(r, "AU")
        ac = g.get(r, "AR")
        dex = g.get(r, "AS", keep_zero=True)      # 重甲「敏捷加值 0」是有意义的值，保留
        weight = g.get(r, "AT")
        if not name:
            recovered = rec.name_for(r, "AQ")
            if recovered:
                name = recovered
                healed.append((r, recovered))
            elif ac or dex or weight or note:
                rep["skipped_rows"].append(
                    (f"AQ–AU 第 {r} 行", g.raw(r, "AQ") or "(空)",
                     "无护甲名称（AQ='-'，只剩基础 AC 10 / 敏捷加值），无法命名"))
                continue
            else:
                continue
        fields = {"护甲等级": ac, "敏捷加值": dex, "重量": weight, "属性": note}
        if group:
            fields["类型"] = group
        if "隐匿劣势" in note:
            fields["隐匿劣势"] = "劣势"
        m = re.search(r"力量需求\s*(\d+)", note)
        if m:
            fields["力量需求"] = m.group(1)
        # 无分组的前半段（法师护甲 / 大法师法袍 / 职业与种族的无甲防御）不是护甲物品，只是 AC 来源
        b.add(name, "护甲" if group else "其它", r, text=note, fields=fields,
              tags=[group] if group else ["AC来源"])
        rows.append(r)
    rep["blocks"].append({
        "name": "护甲表", "cols": "AQ–AU", "header_row": 20,
        "header": "护甲名称 | AC | 敏捷加值 | 重量 | 属性",
        "rows": f"{min(rows)}–{max(rows)}（跳过分隔行 {dividers}）", "count": len(rows),
        "note": "第 21 行 AQ='-'（无名称，基础 AC 10）跳过；第 22–35 行是无甲防御/法术/魔法物品的 AC 来源，"
                "归入 category=其它 + tags=['AC来源']；"
                + (f"行 {[r for r, _ in healed]} 的名称由 xlsx 公式恢复" if healed else "无名称被恢复"),
    })
    for r, name in healed:
        rep["recovered_names"].append((f"AQ{r}", name))


def parse_weapons(g: Grid, b: Builder, rep: dict, rec: NameRecovery) -> None:
    """AI–AP 块：名称 | 伤害 | 伤害类型 | 词条 | 精通 | 重量 | 价格 | 属性（表头第 1 行）。
    AH 列是分组标题（简易/军用 × 近战/远程），按行向下沿用；行 78–80 是「自定义武器」空模板。"""
    labels = {}
    for r in range(2, g.rows + 1):
        lab = g.get(r, "AH")
        if lab:
            labels[r] = re.match(r"^[^\sA-Za-z]+", lab).group(0)   # 去掉尾部英文
    label_rows = sorted(labels)
    first_label = labels[label_rows[0]] if label_rows else ""
    rows = []
    for r in range(2, g.rows + 1):
        name = g.get(r, "AI") or rec.name_for(r, "AI")
        damage = g.get(r, "AJ")
        if not name:
            if damage:
                rep["skipped_rows"].append((f"AI–AP 第 {r} 行", "(空)", "有伤害数据但无武器名称"))
            continue
        if not damage:
            rep["skipped_rows"].append((f"AI–AP 第 {r} 行", name, "无伤害数据（自定义武器空模板/占位行）"))
            continue
        prior = [lr for lr in label_rows if lr <= r]
        wtype = labels[prior[-1]] if prior else first_label
        b.add(name, "武器", r, text="",
              fields={"类型": wtype, "伤害": damage, "伤害类型": g.get(r, "AK"),
                      "属性": g.get(r, "AL"), "精通": g.get(r, "AM"),
                      "重量": g.get(r, "AN"), "价格": g.get(r, "AO"),
                      "检定属性": g.get(r, "AP")},
              tags=[wtype] if wtype else [])
        rows.append(r)
    rep["blocks"].append({
        "name": "武器表", "cols": "AI–AP", "header_row": 1,
        "header": "名称 | 伤害 | 伤害类型 | 词条 | 精通 | 重量 | 价格 | 属性",
        "rows": f"{min(rows)}–{max(rows)}", "count": len(rows),
        "note": "AH 列分组标题：" + "、".join(f"第 {r} 行→{labels[r]}" for r in label_rows)
                + "；行 3（徒手打击）在第一个标题行之上，按第一个标题归属",
    })


# --------------------------------------------------------------------------- #
# 跳过的块（明确记录，不猜结构）
# --------------------------------------------------------------------------- #
def audit_skipped(g: Grid, rep: dict, main_names: set[str]) -> None:
    eh = [(r, g.get(r, "E")) for r in range(2, g.rows + 1) if g.get(r, "E")]
    dup = [n for _, n in eh if n in main_names]
    filler = [r for r in range(2, g.rows + 1) if g.raw(r, "E").strip() == "0"]
    rep["skipped_blocks"].append({
        "cols": "E–H", "header_row": 1, "header": "物品 | 特性 | 重量 | 价格",
        "reason": f"主表 A–D 的镜像副本（重量/价格列序相反）：有名称的 {len(eh)} 行里 {len(dup)} 行与主表同名"
                  f"（强酸还重复了 2 次）；第 {min(filler)}–{max(filler)} 行 E/F/G 全是 '0'；"
                  "第 13 行 E/F/G = 背包1/背包2/次元袋 是残留。为避免重复不收。",
    })
    rep["skipped_blocks"].append({
        "cols": "I–J", "header_row": None, "header": "（无表头）",
        "reason": f"只有第 2 行一格数据（{g.get(2, 'I')!r} + 描述），与 A–D 第 19 行重复。",
    })
    rep["skipped_blocks"].append({
        "cols": "R–W", "header_row": 1, "header": "工具 | 属性 | 操作 | 制造 | 重量 | 价格",
        "reason": "只有表头，没有数据：R2:W11 全是形如 "
                  "=IF(ISBLANK(背包!B16),,INDEX($L$2:$Q$42,MATCH(背包!B16&\"*\",...))) 的查表公式，"
                  "指向「背包」工作表，当前缓存结果全为空。属于角色自身的背包查表视图，不是物品目录。",
    })
    zz = [(r, g.get(r, "AA")) for r in range(2, g.rows + 1) if g.get(r, "AA")]
    rep["skipped_blocks"].append({
        "cols": "Z–AD", "header_row": 1,
        "header": "弹药重量(Z) | 集束重量(AB) | 集束数量(AC) | 总重量→(AD)",
        "reason": "弹药汇总小表：名称列 Z 只有 箭/矢/弹/投/它 这类残字（与 AA 列 '箭arrows(20)'、"
                  "'弩矢crossbowbolts(20)' 对不上，疑为公式残留），总重量列 AD 全为占位 0；"
                  f"AA 列的 {len(zz)} 项（箭arrows(20)/弩矢crossbowbolts(20)/子弹/"
                  "投石索弹丸slingbullets(20)/吹矢blowgunneedles(50)）与主表 A–D 第 4–8 行"
                  "（箭矢/弩矢/枪械子弹/投石索子弹/吹矢）是同一批弹药。为免产生「箭矢」与「箭」这类重复条目，"
                  "本块不作为词条来源；弹药词条由主表提供（集束重量/集束数量因此未收录）。",
    })
    rep["skipped_blocks"].append({
        "cols": "AQ–AZ 第 2–19 行", "header_row": 1,
        "header": "装备名称 | 装备特性 | 伤害 | 伤害类型 | 精通 | 命中检定 | 调整值 | 伤害加值 | 重量",
        "reason": "角色卡自身的攻击计算区（匕首 命中+3 / 伤害+1，其余为 '0'）：AQ/AR… 全是 "
                  "VLOOKUP(AQ,$AI$2:$AP$80,…) 加 主要! 表引用，第 11 行是 武器命中/武器伤害 标签，"
                  "第 18–19 行是 盾牌/0 残留。不是物品目录，不收。",
    })
    rep["skipped_blocks"].append({
        "cols": "AV–AZ、BA", "header_row": 20,
        "header": "护甲类型 | AC | 敏捷加值 | 重量 | 属性",
        "reason": "查表暂存区：AV–AZ 只有第 21 行有值（兽皮甲 12/1/12），与 AQ–AU 第 42 行 兽皮甲 完全一致；"
                  "BA 列是同一份护甲名单整体下移一行的副本（第 28 行起与 AQ 错位），无新增数据。",
    })
    rep["skipped_blocks"].append({
        "cols": "K、W、X、AC/AD/AE(第 9 行起)、AF、AG、BB、BE",
        "header_row": None, "header": "（辅助/下拉列）",
        "reason": "K = L 列工具名单上移一行；W(护甲/武器/奇物/盾牌)、X(装备同调数量 0)、"
                  "AC(武器精通+属性)、AD(身体部位)、AE(体型/护甲类型/伤害类型) 是数据验证下拉清单；"
                  "AG 是 AI 的镜像列；AF 只有 '-/√/OK'，BB/BE 只有零散数字。均非词条。",
    })


# --------------------------------------------------------------------------- #
def main() -> int:
    dry = "--dry-run" in sys.argv
    use_xlsx = "--no-xlsx" not in sys.argv

    g = Grid(GRID_PATH)
    rec = NameRecovery(XLSX_PATH, SHEET_NAME) if use_xlsx else NameRecovery(Path("__none__"), SHEET_NAME)
    b = Builder()
    rep = {"blocks": [], "skipped_blocks": [], "skipped_rows": [], "recovered_names": []}

    main_names = {g.get(r, "A") for r in range(2, g.rows + 1) if g.get(r, "A")}

    parse_main(g, b, rep, rec)
    parse_tools(g, b, rep, rec)
    parse_armor(g, b, rep, rec)
    parse_weapons(g, b, rep, rec)
    audit_skipped(g, rep, main_names)

    # 校验
    ids = [e["id"] for e in b.entries]
    assert len(ids) == len(set(ids)), "id 不唯一"
    assert all(e["name"] for e in b.entries), "存在空名称"
    assert all(e["cardRef"]["sheet"] == SHEET_NAME for e in b.entries), "sheet 名不对"
    assert all(e["cardRef"]["row"] > 1 for e in b.entries), "cardRef 行号异常"
    bad = [v for e in b.entries for v in list(e["fields"].values()) + [e["text"]]
           if is_placeholder(v) and v.strip() not in ("", "0")]
    assert not bad, f"字段里混入占位符: {bad[:5]}"

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    if not dry:
        OUT_PATH.write_text(json.dumps(b.entries, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    # ---------------- 报告 ----------------
    print(f"网格：{GRID_PATH}  sheet={g.sheet} rows={g.rows} cols={g.cols} cells={len(g.cells)}")
    print(f"xlsx：{rec.note}")
    print(f"输出：{'(dry-run 未写)' if dry else OUT_PATH}")

    print(f"\n词条总数：{len(b.entries)}")
    dist: dict[str, int] = {}
    for e in b.entries:
        dist[e["category"]] = dist.get(e["category"], 0) + 1
    for k, v in sorted(dist.items(), key=lambda kv: -kv[1]):
        print(f"  {k}: {v}")

    print("\n== 识别出的块 ==")
    print(f"{'列区间':<12} {'表头行':<7} {'词条数':<7} {'表头':<50} 块名")
    for blk in rep["blocks"]:
        print(f"{blk['cols']:<12} {blk['header_row']:<7} {blk['count']:<7} {blk['header']:<50} {blk['name']}")
    for blk in rep["blocks"]:
        if blk.get("note"):
            print(f"  · [{blk['name']}] {blk['note']}")

    print("\n== 跳过的块 ==")
    for blk in rep["skipped_blocks"]:
        print(f"  · {blk['cols']}（表头行 {blk['header_row']}）: {blk['reason']}")

    if rep["recovered_names"]:
        print("\n== 由 xlsx 公式恢复的名称 ==")
        for cell, name in rep["recovered_names"]:
            print(f"  · {cell} → {name!r}（网格导出缺格，公式里的中文字面量）")

    if rep["skipped_rows"]:
        print("\n== 块内跳过的行 ==")
        for where, name, why in rep["skipped_rows"]:
            print(f"  · {where} {name!r}: {why}")

    if b.collisions:
        print("\n== 重名（id 加 #n 后缀）==")
        for name, first, again in b.collisions:
            print(f"  · {name!r}: 第 {first} 行 → 第 {again} 行")
    else:
        print("\n无重名。")

    print(f"\n网格中出现的占位符取值：{sorted({v for v in g.cells.values() if is_placeholder(v)})}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
